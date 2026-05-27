from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


FONT_PATH = Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf")
PAGE = landscape((7.5 * inch, 13.333 * inch))
BLUE = colors.HexColor("#2F6B8F")
RED = colors.HexColor("#C46243")
GREEN = colors.HexColor("#88A868")
INK = colors.HexColor("#1F2933")
MUTED = colors.HexColor("#52606D")
LIGHT = colors.HexColor("#EEF3F6")


def font_name() -> str:
    if FONT_PATH.exists() and "ArialUnicode" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("ArialUnicode", str(FONT_PATH)))
    return "ArialUnicode" if FONT_PATH.exists() else "Helvetica"


def wrap(text: str, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    line = ""
    for word in words:
        candidate = f"{line} {word}".strip()
        if len(candidate) > width and line:
            lines.append(line)
            line = word
        else:
            line = candidate
    if line:
        lines.append(line)
    return lines


def draw_text(c: canvas.Canvas, text: str, x: float, y: float, size: int = 18, color=INK, max_width: int = 80, leading: int | None = None) -> float:
    leading = leading or int(size * 1.35)
    c.setFillColor(color)
    c.setFont(font_name(), size)
    for line in wrap(text, max_width):
        c.drawString(x, y, line)
        y -= leading
    return y


def draw_bullets(c: canvas.Canvas, bullets: list[str], x: float, y: float, size: int = 17, max_width: int = 78) -> float:
    c.setFont(font_name(), size)
    c.setFillColor(INK)
    for bullet in bullets:
        lines = wrap(bullet, max_width)
        c.drawString(x, y, "•")
        c.drawString(x + 0.25 * inch, y, lines[0])
        y -= size * 1.35
        for line in lines[1:]:
            c.drawString(x + 0.25 * inch, y, line)
            y -= size * 1.35
        y -= 0.08 * inch
    return y


def slide_header(c: canvas.Canvas, title: str, subtitle: str | None = None) -> None:
    c.setFillColor(BLUE)
    c.rect(0, PAGE[1] - 0.18 * inch, PAGE[0], 0.18 * inch, fill=1, stroke=0)
    draw_text(c, title, 0.65 * inch, PAGE[1] - 0.72 * inch, 26, INK, 70)
    if subtitle:
        draw_text(c, subtitle, 0.65 * inch, PAGE[1] - 1.08 * inch, 13, MUTED, 110)


def add_metric(c: canvas.Canvas, label: str, value: str, x: float, y: float, color=BLUE) -> None:
    c.setFillColor(LIGHT)
    c.roundRect(x, y, 2.65 * inch, 0.86 * inch, 6, fill=1, stroke=0)
    c.setFillColor(color)
    c.setFont(font_name(), 22)
    c.drawString(x + 0.18 * inch, y + 0.45 * inch, value)
    c.setFillColor(MUTED)
    c.setFont(font_name(), 10)
    c.drawString(x + 0.18 * inch, y + 0.18 * inch, label)


def add_image(c: canvas.Canvas, path: Path, x: float, y: float, w: float, h: float) -> None:
    if path.exists():
        c.drawImage(str(path), x, y, width=w, height=h, preserveAspectRatio=True, anchor="c")


def build_pdf(output_path: Path, metrics: dict, root_cause: pd.DataFrame, figures_dir: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(output_path), pagesize=PAGE)
    c.setTitle("GContest Fraud Anomaly Detection Slide Deck")
    prevention = metrics.get("prevention_impact", {})
    model_metrics = metrics.get("supervised_model_metrics", {})

    slide_header(c, "Fraud & Anomaly Detection", "Cause-first risk framework using real G'Contest banking data")
    add_metric(c, "Transactions scored", f"{metrics['row_counts']['transactions_scored']:,}", 0.65 * inch, 4.85 * inch)
    add_metric(c, "High/Critical cases", f"{metrics['review_queue']['high_or_critical_transactions']:,}", 3.55 * inch, 4.85 * inch, RED)
    add_metric(c, "Customers impacted", f"{metrics['review_queue']['customers_with_high_or_critical']:,}", 6.45 * inch, 4.85 * inch, GREEN)
    draw_bullets(
        c,
        [
            "No synthetic data and no fake fraud labels.",
            "Output is a prevention queue with explainable risk reasons.",
            "Framework aligns with risk-based banking controls: step-up authentication, manual review, AML escalation.",
            "Optimization priority: catch more fraud-risk cases first, while reporting false-positive rate for customer experience control.",
        ],
        0.8 * inch,
        4.15 * inch,
        18,
        84,
    )
    c.showPage()

    slide_header(c, "Why Cause-First", "BGK note: identify root causes before introducing the model")
    draw_bullets(
        c,
        [
            "Account takeover / identity compromise: new device, new IP, night access, new beneficiary, late-stage digital activity.",
            "Unauthorized transfer / capital outflow: outside-bank transfer, amount above customer baseline, daily burst, cash-out versus CASA.",
            "AML network / mule-account pattern: shared IP/device/beneficiary, round high-value transfers, repeated external transfers.",
        ],
        0.8 * inch,
        5.55 * inch,
        17,
        92,
    )
    c.showPage()

    slide_header(c, "4-Phase Framework", "The final system follows the requested 11-step prevention flow")
    steps = [
        ("1. Foundation", "Clean data, create 4 baseline groups, package Customer 360 with rolling 30/60/90 days"),
        ("2. EDA + Labels", "Insight charts, IQR thresholding, dynamic root-cause rules, weak labels"),
        ("3. ML + Matrix", "RandomForest learns weak labels; Rule+ML matrix converts signals into bank actions"),
        ("4. Operations", "Streamlit dashboard, protected amount, SHAP xAI and human-readable reason codes"),
    ]
    x = 0.65 * inch
    for idx, (title, body) in enumerate(steps):
        y = 5.0 * inch - idx * 0.95 * inch
        c.setFillColor(LIGHT if idx % 2 == 0 else colors.white)
        c.roundRect(x, y, 11.9 * inch, 0.72 * inch, 5, fill=1, stroke=0)
        draw_text(c, title, x + 0.22 * inch, y + 0.40 * inch, 15, BLUE, 24)
        draw_text(c, body, x + 2.0 * inch, y + 0.40 * inch, 14, INK, 82)
    c.showPage()

    slide_header(c, "Prevention Performance", "Weak-label evaluation for an operational fraud prevention queue")
    add_metric(c, "Recall / prevention coverage", f"{prevention.get('prevention_coverage_against_rule_labels', 0):.1%}", 0.65 * inch, 4.85 * inch, RED)
    add_metric(c, "Precision of Step-up/Block", f"{prevention.get('step_up_or_block_precision_against_rule_labels', 0):.1%}", 3.55 * inch, 4.85 * inch, BLUE)
    add_metric(c, "False-positive rate", f"{model_metrics.get('validation_false_positive_rate_at_high_threshold', 0):.2%}", 6.45 * inch, 4.85 * inch, GREEN)
    draw_bullets(
        c,
        [
            "Metrics are measured against rule-derived weak labels because the contest data has no confirmed fraud labels.",
            "The bank objective is risk minimization: prioritize catching suspected fraud, then tune thresholds against review capacity and customer friction.",
            f"Protected amount by Block/Step-up queue: {prevention.get('protected_amount_block_or_step_up', 0):,.0f} VND.",
            f"High-threshold confusion matrix: {model_metrics.get('validation_confusion_matrix_at_high_threshold', {})}.",
        ],
        0.8 * inch,
        4.0 * inch,
        15,
        98,
    )
    c.showPage()

    slide_header(c, "Customer 360 Baseline", "One customer, one current financial-behavioral profile")
    add_image(c, figures_dir / "rolling_window_baseline.png", 0.75 * inch, 1.05 * inch, 5.6 * inch, 4.4 * inch)
    add_image(c, figures_dir / "customer_360_credit_risk_group.png", 6.75 * inch, 1.05 * inch, 5.6 * inch, 4.4 * inch)
    draw_bullets(
        c,
        [
            "Transactional: amount average/P95/IQR, golden hour, rolling 30/60/90-day activity.",
            "Financial: CASA/TD balance, card utilization, overdue and credit-risk group.",
            "Environmental and behavioral: trusted device/IP, known beneficiary, app activity, night/late-stage activity.",
        ],
        0.85 * inch,
        0.85 * inch,
        13,
        100,
    )
    c.showPage()

    slide_header(c, "Data-Driven Insights", "Charts are selected from actual risk lifts and root-cause interactions")
    insight_bullets = [
        f"{item.get('title')}: {item.get('evidence')}"
        for item in metrics.get("data_driven_insights", [])[:4]
    ]
    draw_bullets(c, insight_bullets, 0.75 * inch, 5.25 * inch, 13, 112)
    c.showPage()

    slide_header(c, "Insight Evidence", "Heatmaps and lift charts replace one-dimensional bar-only EDA")
    add_image(c, figures_dir / "root_cause_hybrid_heatmap.png", 0.55 * inch, 1.1 * inch, 5.9 * inch, 4.3 * inch)
    add_image(c, figures_dir / "iqr_breach_lift.png", 6.75 * inch, 1.1 * inch, 5.7 * inch, 4.3 * inch)
    c.showPage()

    slide_header(c, "Hybrid Decision Matrix", "Rule engine and ML confidence are converted into bank actions")
    add_image(c, figures_dir / "hybrid_decision_matrix.png", 0.8 * inch, 1.0 * inch, 11.7 * inch, 4.8 * inch)
    draw_bullets(
        c,
        [
            "Rule+ML alert: Block/Hold.",
            "Rule-only alert: Step-up/eKYC to reduce false positives.",
            "ML-only alert: Special watchlist for patterns not yet covered by rules.",
            "No alert: Allow while baseline continues to update.",
            "Risk bands are assigned after this matrix so each band maps to a concrete bank control.",
        ],
        0.9 * inch,
        1.0 * inch,
        13,
        100,
    )
    c.showPage()

    slide_header(c, "Temporal Stability", "2019 monthly/quarterly backtest for review-rate stability")
    add_image(c, figures_dir / "monthly_stability_backtest.png", 0.8 * inch, 1.0 * inch, 11.7 * inch, 4.8 * inch)
    draw_bullets(
        c,
        [
            "Only 2019 data is available, so this is a temporal robustness check rather than a crisis-period validation.",
            "The KPI is not confirmed-fraud accuracy; it is weak-label prevention coverage, protected amount, and reason-code consistency over time.",
        ],
        0.9 * inch,
        1.05 * inch,
        14,
        95,
    )
    c.showPage()

    slide_header(c, "Risk Pattern Maps", "Time, network and Customer 360 context explain where review capacity should go")
    add_image(c, figures_dir / "time_risk_heatmap.png", 0.55 * inch, 1.15 * inch, 3.9 * inch, 4.1 * inch)
    add_image(c, figures_dir / "network_exposure_bubble.png", 4.7 * inch, 1.15 * inch, 3.8 * inch, 4.1 * inch)
    add_image(c, figures_dir / "customer360_risk_heatmap.png", 8.85 * inch, 1.15 * inch, 3.7 * inch, 4.1 * inch)
    c.showPage()

    slide_header(c, "xAI Engine", "SHAP explains a tree surrogate of the supervised prevention score")
    add_image(c, figures_dir / "shap_feature_importance.png", 0.7 * inch, 1.05 * inch, 5.8 * inch, 4.5 * inch)
    add_image(c, figures_dir / "shap_summary_beeswarm.png", 6.7 * inch, 1.05 * inch, 5.8 * inch, 4.5 * inch)
    c.showPage()

    slide_header(c, "Operational Recommendations", "Turn model output into bank actions")
    draw_bullets(
        c,
        [
            "Critical: temporary hold, near-real-time manual review, customer verification, step-up authentication.",
            "High: step-up authentication before allowing the transaction to proceed.",
            "AML branch: network escalation rather than single-transaction review only.",
            "Medium: enhanced monitoring; escalate if repeated within the next review window.",
            "Feedback loop: investigator decisions become labels for Precision@K, Recall@K and supervised model tuning.",
        ],
        0.8 * inch,
        5.55 * inch,
        17,
        92,
    )
    c.showPage()

    slide_header(c, "Live Demo", "Streamlit advisor for transaction/customer review")
    draw_bullets(
        c,
        [
            "Run `streamlit run src/demo_app.py`.",
            "Select a top-risk transaction or enter a CUSTOMER_NUMBER.",
            "The advisor returns risk band, root-cause branch, reason codes, recommended action and SHAP evidence.",
            "This is designed for a short judging demo or screen-recorded walkthrough.",
        ],
        0.8 * inch,
        5.55 * inch,
        18,
        88,
    )
    c.save()


def set_text(box, text: str, font_size: int = 18, color: RGBColor | None = None, bold: bool = False) -> None:
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def build_pptx(output_path: Path, metrics: dict, figures_dir: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    prevention = metrics.get("prevention_impact", {})
    model_metrics = metrics.get("supervised_model_metrics", {})
    slides = [
        ("Fraud & Anomaly Detection", "Cause-first framework using real data, no synthetic labels.", []),
        (
            "Cause-First Hypothesis",
            "Three root-cause branches before modeling.",
            [
                "Account takeover / identity compromise",
                "Unauthorized transfer / capital outflow",
                "AML network / mule-account pattern",
            ],
        ),
        (
            "Framework",
            "4 phases, 11 steps: foundation -> EDA/labels -> ML/matrix -> dashboard/xAI.",
            [
                "Phase 1: clean data, four baseline groups, Customer 360 rolling 30/60/90 days",
                "Phase 2: strategic EDA, IQR thresholding, dynamic root-cause rules, weak labels",
                "Phase 3: RandomForest prevention model and Rule + ML action matrix",
                "Phase 4: dashboard, protected amount, SHAP xAI and natural-language reason codes",
            ],
        ),
        (
            "Data Insights",
            "EDA is now driven by risk lift, heatmaps, and interaction charts.",
            [
                "Root-cause x hybrid-decision heatmap",
                "IQR breach lift versus normal baseline",
                "Time, network and Customer 360 risk maps",
            ],
        ),
        (
            "Key Metrics",
            f"{metrics['row_counts']['transactions_scored']:,} transactions scored; {metrics['review_queue']['high_or_critical_transactions']:,} High/Critical.",
            [
                f"Prevention coverage vs weak labels: {prevention.get('prevention_coverage_against_rule_labels', 0):.1%}",
                f"False-positive rate at high threshold: {model_metrics.get('validation_false_positive_rate_at_high_threshold', 0):.2%}",
                "Confusion matrix and protected amount are reported for business impact, not only model accuracy.",
            ],
        ),
    ]
    for title, subtitle, bullets in slides:
        slide = prs.slides.add_slide(blank)
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.16))
        band.fill.solid()
        band.fill.fore_color.rgb = RGBColor(47, 107, 143)
        band.line.fill.background()
        title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.55), Inches(12), Inches(0.55))
        set_text(title_box, title, 30, RGBColor(31, 41, 51), True)
        sub_box = slide.shapes.add_textbox(Inches(0.58), Inches(1.15), Inches(12), Inches(0.45))
        set_text(sub_box, subtitle, 16, RGBColor(82, 96, 109))
        if bullets:
            body = slide.shapes.add_textbox(Inches(0.75), Inches(2.0), Inches(11.7), Inches(4.2))
            tf = body.text_frame
            tf.clear()
            for idx, bullet in enumerate(bullets):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.text = bullet
                p.font.name = "Arial"
                p.font.size = Pt(21)
                p.level = 0
    for title, image in [
        ("Root-Cause x Hybrid Matrix", "root_cause_hybrid_heatmap.png"),
        ("IQR Breach Lift", "iqr_breach_lift.png"),
        ("Time Risk Heatmap", "time_risk_heatmap.png"),
        ("Network Exposure", "network_exposure_bubble.png"),
        ("Customer 360 Risk Map", "customer360_risk_heatmap.png"),
        ("Customer 360 Rolling Baseline", "rolling_window_baseline.png"),
        ("Hybrid Decision Matrix", "hybrid_decision_matrix.png"),
        ("Temporal Stability", "monthly_stability_backtest.png"),
        ("SHAP Explainability", "shap_feature_importance.png"),
    ]:
        slide = prs.slides.add_slide(blank)
        title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.4), Inches(12), Inches(0.5))
        set_text(title_box, title, 28, RGBColor(31, 41, 51), True)
        image_path = figures_dir / image
        if image_path.exists():
            slide.shapes.add_picture(str(image_path), Inches(0.9), Inches(1.2), width=Inches(11.5), height=Inches(5.7))
    prs.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build final slide deck as PDF and PPTX.")
    parser.add_argument("--outputs-dir", type=Path, default=Path("outputs"))
    parser.add_argument("--figures-dir", type=Path, default=Path("outputs/figures"))
    parser.add_argument("--report-dir", type=Path, default=Path("report"))
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    metrics = json.loads((args.outputs_dir / "model_metrics.json").read_text(encoding="utf-8"))
    root_cause = pd.read_csv(args.outputs_dir / "root_cause_summary.csv")
    build_pdf(args.report_dir / "final_slide_deck.pdf", metrics, root_cause, args.figures_dir)
    build_pptx(args.report_dir / "final_slide_deck.pptx", metrics, args.figures_dir)
    print(f"Wrote {(args.report_dir / 'final_slide_deck.pdf').resolve()}")
    print(f"Wrote {(args.report_dir / 'final_slide_deck.pptx').resolve()}")


if __name__ == "__main__":
    main()
